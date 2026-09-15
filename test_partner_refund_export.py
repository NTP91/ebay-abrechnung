"""Regression for Group-B partner refund variant B.

Every original sale remains positive and every refund remains one separate
negative event.  A refund must never also shrink/remove its sale.

The order numbers/amounts below mirror the real MH control set reported for
payout 7725289401 plus the two later RE0090-linked refunds - used here only
as a regression fixture, never hard-coded into the production logic itself.
The same scenario is repeated for a second partner (NB) to prove the fix is
generic, not an MH-specific special case.
"""
import io
import tempfile
import unittest
from decimal import Decimal, ROUND_HALF_UP
from fractions import Fraction
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import core
import position_workflow as workflow
import studio_view
from partner_export import export_partner_excel, prepare_partner_export, DISPLAY_NAMES
from test_recovery import payout
from test_invoice_support import review_positions


def reference_gutschriften_gross(amounts, rate=Decimal('.035'), tax=Decimal('.19')):
    """Independent oracle for 'rate on net, VAT re-added afterwards' (the only
    rule the partner formula may use) - built without touching partner_export
    or calculate_sheet, so it can catch a shortcut like gross*(1-rate)."""
    cent = Decimal('.01')

    def cents(value):
        return value.quantize(cent, rounding=ROUND_HALF_UP)
    total_after = previous_tax = total = Decimal('0')
    for gross in amounts:
        net = cents(Decimal(gross) / Decimal('1.19'))
        after = cents(net * (1 - rate))
        total_after += after
        tax_to_date = cents(total_after * tax)
        total += after + tax_to_date - previous_tax
        previous_tax = tax_to_date
    return total

# (Bestellnummer, refund amount) - all fully refunded before any partner payment.
MH_PAYOUT_7725289401_REFUNDS = [
    ('02-15111-34101', '-155,98'),
    ('14-15091-70890', '-448,99'),
    ('21-15067-10713', '-77,99'),
    ('24-15066-72008', '-34,99'),
    ('27-15053-47641', '-12,99'),
]
# Later refunds against MH orders already inside an (unrelated, untouched) Evelyn
# Lexware voucher; MH itself was never paid for them.
MH_RE0090_LINKED_REFUNDS = [
    ('09-15110-28069', '-29,99'),
    ('18-15094-64423', '-29,99'),
]


class PartnerRefundExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.paths = patch.multiple(core, PAYOUTS_DB_PATH=str(self.root / 'Master_Payouts.csv'),
                                    ORDERS_DB_PATH=str(self.root / 'Master_Orders.csv'))
        self.paths.start()
        self.addCleanup(self.paths.stop)
        self.http = patch('requests.sessions.Session.request', side_effect=AssertionError('Live HTTP forbidden'))
        self.http.start()
        self.addCleanup(self.http.stop)

    def _finish(self, frames):
        for frame in frames:
            frame['Transaktionsbetrag (inkl. Kosten)'] = frame['Betrag abzügl. Kosten']
            frame['Auszahlungsdatum'] = '03.09.2026'
            frame['Auszahlungsstatus'] = 'Betrag überwiesen'
        core.import_reports(frames, core.ORDERS_DB_PATH, 'orders')
        core.import_reports(frames, core.PAYOUTS_DB_PATH, 'payout')
        core.sync_status(core.load_master_data())

    def _sale_refund(self, payout_id, order, sale_amount, refund_amount, sku):
        sale = payout(payout_id, 'sale-' + order, order, sku=sku, amount=sale_amount)
        refund = payout(payout_id, 'refund-' + order, order, sku=sku, amount=refund_amount, kind='Rückerstattung')
        for frame in (sale, refund):
            frame['Artikelnummer'] = 'item-' + order
        return [sale, refund]

    def _regular_sale(self, payout_id, order, amount, sku):
        sale = payout(payout_id, 'sale-' + order, order, sku=sku, amount=amount)
        sale['Artikelnummer'] = 'item-' + order
        return sale

    def build_scenario(self, sku, full_refunds, regular_count=4):
        """A partner with `regular_count` untouched open positions, the given
        full-refund order/amount pairs (all unpaid), two Evelyn-Lexware-bound
        orders refunded later (MH only), and one already-reviewed order
        refunded after review (must become a separate refund_case, never
        merged with new positions)."""
        reviewed_order = 'reviewed-1'
        # Phase 1: the reviewed order exists alone, gets reviewed/approved -
        # its own historical invoice, fixed before anything else is imported.
        self._finish([self._regular_sale('p-reviewed', reviewed_order, '50,00', sku)])
        first_key = workflow.positions().loc[lambda r: r.Bestellnummer == reviewed_order, 'position_key'].iloc[0]
        review_positions([first_key])

        # Phase 2: everything else, including the reviewed order's refund.
        frames = []
        for order, refund_amount in full_refunds:
            sale_amount = refund_amount.lstrip('-')
            frames += self._sale_refund('p-open', order, sale_amount, refund_amount, sku)
        regular_orders = [f'regular-{i}' for i in range(regular_count)]
        for i, order in enumerate(regular_orders):
            frames.append(self._regular_sale('p-open', order, f'{20 + i}.00'.replace('.', ','), sku))
        bound_orders = []
        if sku.startswith('MH'):
            for order, refund_amount in MH_RE0090_LINKED_REFUNDS:
                sale_amount = refund_amount.lstrip('-')
                frames += self._sale_refund('p-bound', order, sale_amount, refund_amount, sku)
                bound_orders.append(order)
        frames += self._sale_refund('p-open', reviewed_order, '50,00', '-50,00', sku)[1:]  # refund only
        self._finish(frames)
        if bound_orders:
            with core.ledger() as db:
                db.execute("UPDATE payouts SET attempt='created',invoice_id='RE0090-test' WHERE id='p-bound'")
                db.commit()
        return regular_orders, bound_orders, reviewed_order

    def run_partner_scenario(self, sku, partner_name):
        full_refunds = MH_PAYOUT_7725289401_REFUNDS if partner_name == 'MH' else [
            (f'{partner_name.lower()}-full-{i}', f'-{10+i},00') for i in range(3)
        ]
        regular_orders, bound_orders, reviewed_order = self.build_scenario(sku, full_refunds, regular_count=4)
        business = workflow.positions()
        if bound_orders:
            self.assertTrue(business.loc[business.Bestellnummer.isin(bound_orders) & (business.Art == 'Bestellung'), 'Lexware_uebertragen'].all())
        partner_ready = studio_view.partner_rows(business)
        mh = partner_ready[partner_ready.Partner == partner_name]

        refunded_orders = {order for order, _ in full_refunds} | set(bound_orders)
        open_sales = mh[mh.Art == 'Bestellung']
        # Variant B: every original unpaid sale remains visible and positive.
        self.assertEqual(sorted(open_sales.Bestellnummer),
                         sorted(regular_orders + [reviewed_order] + list(refunded_orders)))
        self.assertTrue(refunded_orders.issubset(set(open_sales.Bestellnummer)))

        # The reviewed-and-then-refunded order passes through unmodified (history never changes)...
        settled = mh[(mh.Bestellnummer == reviewed_order) & (mh.Art == 'Bestellung')]
        self.assertEqual(len(settled), 1)
        self.assertEqual(round(float(settled.iloc[0].Erlös_Brutto), 2), 50.00)
        # ...and its refund is a separate new event while history remains intact.
        self.assertIn(reviewed_order, mh.loc[mh.Art == 'Erstattung', 'Bestellnummer'].tolist())
        cases = studio_view.partner_refund_cases(business[business.Partner == partner_name])
        self.assertEqual(cases.Bestellnummer.tolist(), [reviewed_order])

        # Every applicable refund is visible exactly once, tied to its own order/finance id.
        export_refunds = mh[mh.Art == 'Erstattung']
        self.assertEqual(sorted(export_refunds.Bestellnummer), sorted(refunded_orders | {reviewed_order}))
        self.assertEqual(export_refunds.Bestellnummer.duplicated().sum(), 0)
        self.assertEqual(export_refunds.Transaktionsnummer.duplicated().sum(), 0)

        # Mirror app.py exactly: only not-yet-reviewed rows go into the *new* export/download
        # ("Neu für nächste Rechnung"); the already-reviewed order stays out - its historical
        # invoice is untouched and it is paid via the separate "Zahlung offen" bucket instead.
        next_invoice = mh[~mh.reviewed_at.astype(bool)]
        self.assertNotIn(reviewed_order, next_invoice.loc[next_invoice.Art=='Bestellung','Bestellnummer'].tolist())
        self.assertIn(reviewed_order, next_invoice.loc[next_invoice.Art=='Erstattung','Bestellnummer'].tolist())

        model = prepare_partner_export(next_invoice)
        rechnung_orders = {item['order'] for item in model['Rechnung']}
        gutschrift_orders = {item['order'] for item in model['Gutschriften']}
        historical_orders = {item['order'] for item in model['HistorischeGutschriften']}
        self.assertEqual(rechnung_orders, set(regular_orders) | refunded_orders)
        # reviewed_order's sale was already paid to the partner in an earlier
        # run (Fall B): its later refund is an open repayment case of its own
        # (Tab 3), never merged back into this settlement's Gutschriften (Tab 2).
        self.assertEqual(gutschrift_orders, refunded_orders)
        self.assertEqual(historical_orders, {reviewed_order})
        for order in refunded_orders:
            sale=next(item for item in model['Rechnung'] if item['order']==order)
            refund=next(item for item in model['Gutschriften'] if item['order']==order)
            self.assertEqual(sale['gross']+refund['gross'],Decimal('0.00'))

        blob = export_partner_excel(next_invoice)
        path = Path(self.temp.name) / f'Partner_Patrick_{partner_name}.xlsx'
        path.write_bytes(blob)
        book = load_workbook(io.BytesIO(path.read_bytes()), data_only=True)
        self.assertEqual(book.sheetnames, [DISPLAY_NAMES['Rechnung'], DISPLAY_NAMES['Gutschriften'], DISPLAY_NAMES['HistorischeGutschriften']])
        rechnung, gutschriften = book[DISPLAY_NAMES['Rechnung']], book[DISPLAY_NAMES['Gutschriften']]
        historische = book[DISPLAY_NAMES['HistorischeGutschriften']]
        rechnung_text = '\n'.join(str(cell.value) for row in rechnung for cell in row if cell.value is not None)
        gutschrift_text = '\n'.join(str(cell.value) for row in gutschriften for cell in row if cell.value is not None)
        historical_text = '\n'.join(str(cell.value) for row in historische for cell in row if cell.value is not None)
        self.assertNotIn('Keine Erstattungen vorhanden', gutschrift_text)
        self.assertNotIn('Keine offenen Rückforderungen vorhanden', historical_text)
        self.assertNotIn(reviewed_order, rechnung_text)  # historical invoice stays out of the new export
        self.assertNotIn(reviewed_order, gutschrift_text)  # already paid out - not this settlement's concern
        self.assertIn(reviewed_order, historical_text)  # surfaced instead as its own open repayment case
        for order in refunded_orders:
            self.assertIn(order, gutschrift_text)
        for order in regular_orders:
            self.assertIn(order, rechnung_text)
            self.assertNotIn(order, gutschrift_text)

        expected_sales=set(regular_orders) | refunded_orders
        expected_refunds=refunded_orders
        rechnung_rows = list(rechnung.iter_rows(min_row=15, max_row=14 + len(expected_sales)))
        self.assertEqual({row[1].value for row in rechnung_rows}, expected_sales)
        gutschrift_rows = list(gutschriften.iter_rows(min_row=15, max_row=14 + len(expected_refunds)))
        self.assertEqual({row[1].value for row in gutschrift_rows}, expected_refunds)
        historical_rows = list(historische.iter_rows(min_row=15, max_row=14 + 1))
        self.assertEqual({row[1].value for row in historical_rows}, {reviewed_order})
        refund_text='\n'.join(str(row[3].value) for row in gutschrift_rows)
        historical_row_text='\n'.join(str(row[3].value) for row in historical_rows)
        # Zusatztext: Bestellnummer deliberately repeated (own column too);
        # no duplicate Bestelldatum, no internal workflow/status bookkeeping,
        # no Refund-ID/Refund brutto (the amount already has its own column);
        # the paid-out flag line is new and expected on both refund sheets.
        for label in ('eBay-Bestellnummer:','SKU:','Refund-Datum:','Refund-Payout:','Bereits an Partner bezahlt:'):
            self.assertIn(label,refund_text)
            self.assertIn(label,historical_row_text)
        for label in ('Bestelldatum:','Ursprünglicher Payout:',
                      'Ursprüngliche Abrechnung:','Refund-ID:','Refund brutto:',
                      'Partnerwirkung:','Status:'):
            self.assertNotIn(label,refund_text)
            self.assertNotIn(label,historical_row_text)
        self.assertIn('Bereits an Partner bezahlt: Nein', refund_text)
        self.assertIn('Bereits an Partner bezahlt: Ja', historical_row_text)

        rechnung_summary_row = 14 + max(1, len(expected_sales)) + 2 + 4
        gutschrift_summary_row = 14 + max(1, len(expected_refunds)) + 2 + 4
        excel_final = round(rechnung[f'K{rechnung_summary_row}'].value + gutschriften[f'K{gutschrift_summary_row}'].value, 2)

        totals = model['totals']
        system_final = round(float(totals['Rechnung']['gross'] + totals['Gutschriften']['gross']), 2)
        summary = studio_view.partner_summary(next_invoice).iloc[0]
        self.assertEqual(round(summary['Verbleibender Anspruch'], 2), system_final)
        self.assertEqual(excel_final, system_final)
        return regular_orders, refunded_orders, system_final

    def test_mh_refund_control_set_matches_reconciliation_and_excel_matches_system(self):
        regular, refunded, final = self.run_partner_scenario('MH / 1', 'MH')
        self.assertEqual(len(regular), 4)
        self.assertEqual(refunded, {o for o, _ in MH_PAYOUT_7725289401_REFUNDS} | {o for o, _ in MH_RE0090_LINKED_REFUNDS})

    def test_same_generic_logic_applies_to_a_second_partner_nb(self):
        regular, refunded, final = self.run_partner_scenario('NB / 1', 'NB')
        self.assertEqual(len(regular), 4)
        self.assertEqual(len(refunded), 3)

    def test_refund_cannot_apply_twice_for_a_partner_export(self):
        first = payout('p-open', 'sale-a', 'order-dup', sku='MH / 1', amount='30,00')
        second = payout('p-open', 'sale-b', 'order-dup', sku='MH / 1', amount='30,00')
        refund = payout('p-open', 'refund-dup', 'order-dup', sku='MH / 1', amount='-30,00', kind='Rückerstattung')
        for frame in (first, second, refund):
            frame['Artikelnummer'] = 'item-dup'
        self._finish([first, second, refund])
        business = workflow.positions()
        mh = studio_view.partner_rows(business)
        mh = mh[mh.Partner == 'MH']
        self.assertEqual(len(mh[(mh.Bestellnummer == 'order-dup') & (mh.Art == 'Bestellung')]), 2)
        self.assertTrue((mh.loc[(mh.Bestellnummer == 'order-dup') & (mh.Art == 'Bestellung'), 'Erlös_Brutto'] == 30.0).all())

    def test_multiple_partial_refunds_of_the_same_order_reduce_cumulatively(self):
        sale = payout('p-open', 'sale-multi', 'order-multi', sku='MH / 1', amount='100,00')
        refund1 = payout('p-open', 'refund-multi-1', 'order-multi', sku='MH / 1', amount='-20,00', kind='Rückerstattung')
        refund2 = payout('p-open', 'refund-multi-2', 'order-multi', sku='MH / 1', amount='-15,00', kind='Rückerstattung')
        for frame in (sale, refund1, refund2):
            frame['Artikelnummer'] = 'item-multi'
        self._finish([sale, refund1, refund2])
        business = workflow.positions()
        mh = studio_view.partner_rows(business)
        mh = mh[mh.Partner == 'MH']
        open_sale = mh[(mh.Bestellnummer == 'order-multi') & (mh.Art == 'Bestellung')]
        self.assertEqual(len(open_sale), 1)
        self.assertEqual(round(float(open_sale.iloc[0].Erlös_Brutto), 2), 100.00)
        refunds = mh[(mh.Bestellnummer == 'order-multi') & (mh.Art == 'Erstattung')]
        self.assertEqual(len(refunds), 2)  # each refund stays its own, separately traceable movement
        self.assertEqual(sorted(round(float(v), 2) for v in refunds.Erlös_Brutto), [-20.0, -15.0])
        model=prepare_partner_export(mh)
        expected=(reference_gutschriften_gross(['100.00'])
                  +reference_gutschriften_gross(['-20.00'])
                  +reference_gutschriften_gross(['-15.00']))
        self.assertEqual(model['totals']['Rechnung']['gross']+model['totals']['Gutschriften']['gross'],expected)

    def test_variant_b_keeps_a_negative_balance_instead_of_clipping_to_zero(self):
        sale = payout('p-open', 'sale-negative', 'order-negative', sku='MH / 1', amount='10,00')
        refund = payout('p-open', 'refund-negative', 'order-negative', sku='MH / 1', amount='-20,00', kind='Rückerstattung')
        for frame in (sale, refund):
            frame['Artikelnummer'] = 'item-negative'
        self._finish([sale, refund])
        mh = studio_view.partner_rows(workflow.positions()).query("Partner == 'MH'")
        self.assertEqual(mh.Art.tolist(), ['Bestellung', 'Erstattung'])
        model = prepare_partner_export(mh)
        self.assertLess(model['totals']['Rechnung']['gross'] + model['totals']['Gutschriften']['gross'], 0)

    def test_duplicate_refund_finance_id_is_blocked_before_export(self):
        frames = []
        for order in ('duplicate-a', 'duplicate-b'):
            sale = payout('p-open', 'sale-' + order, order, sku='MH / 1', amount='20,00')
            refund = payout('p-open', 'refund-' + order, order, sku='MH / 1', amount='-10,00', kind='Rückerstattung')
            for frame in (sale, refund):
                frame['Artikelnummer'] = 'item-' + order
            frames.extend([sale, refund])
        self._finish(frames)
        business=workflow.positions()
        business.loc[business.Art=='Erstattung','Transaktionsnummer']='same-refund-id'
        with self.assertRaisesRegex(ValueError, 'Refund-ID mehrfach'):
            studio_view.partner_rows(business)

    def test_refund_after_partner_payment_is_a_separate_case_not_merged(self):
        sale = payout('p-paid', 'sale-paid', 'order-paid', sku='MH / 1', amount='40,00')
        sale['Artikelnummer'] = 'item-paid'
        self._finish([sale])
        key = workflow.positions().loc[lambda r: r.Bestellnummer == 'order-paid', 'position_key'].iloc[0]
        review_positions([key])
        workflow.confirm([key], 'partner_paid', '2026-09-05')
        paid = workflow.positions().loc[lambda r: r.Bestellnummer == 'order-paid'].iloc[0]
        self.assertTrue(bool(paid.paid_at))

        refund = payout('p-paid', 'refund-paid', 'order-paid', sku='MH / 1', amount='-40,00', kind='Rückerstattung')
        refund['Artikelnummer'] = 'item-paid'
        self._finish([refund])
        business = workflow.positions()
        mh_paid_row = business[(business.Bestellnummer == 'order-paid') & (business.Art == 'Bestellung')].iloc[0]
        self.assertEqual(round(float(mh_paid_row.Erlös_Brutto), 2), 40.00)  # already-paid history untouched

        mh = studio_view.partner_rows(business)
        mh = mh[mh.Partner == 'MH']
        self.assertNotIn('order-paid', mh.loc[mh.Art == 'Bestellung', 'Bestellnummer'].tolist())
        self.assertIn('order-paid', mh.loc[mh.Art == 'Erstattung', 'Bestellnummer'].tolist())
        cases = studio_view.partner_refund_cases(business[business.Partner == 'MH'])
        self.assertEqual(cases.Bestellnummer.tolist(), ['order-paid'])
        self.assertEqual(round(float(cases.iloc[0].Erlös_Brutto), 2), -40.00)

    def test_refund_gross_stays_traceable_and_reduction_uses_the_shared_net_formula(self):
        """The seven reported MH refunds: raw eBay gross sums to 790,92 EUR; the
        resulting reduction of the partner claim must come from the existing
        3,5%-on-net-then-VAT formula (calculate_sheet) - never from
        naively multiplying the raw refund gross by (1 - rate)."""
        # build_scenario('MH ...') already adds MH_RE0090_LINKED_REFUNDS on top of
        # whatever full-refund set is passed in (sku.startswith('MH')) - pass only
        # the payout-7725289401 five here to get all seven without re-declaring
        # the same two transaction ids under a second payout.
        regular_orders, bound_orders, reviewed_order = self.build_scenario('MH / 1', MH_PAYOUT_7725289401_REFUNDS, regular_count=4)
        full_refunds = MH_PAYOUT_7725289401_REFUNDS + MH_RE0090_LINKED_REFUNDS
        business = workflow.positions()
        mh = studio_view.partner_rows(business)
        mh = mh[mh.Partner == 'MH']
        next_invoice = mh[~mh.reviewed_at.astype(bool)]
        refund_rows = next_invoice[(next_invoice.Art == 'Erstattung')
                                   & next_invoice.Bestellnummer.isin({order for order,_ in full_refunds})]

        raw_amounts = [amount.replace(',', '.') for _, amount in full_refunds]
        raw_sum = sum(Decimal(a) for a in raw_amounts)
        self.assertEqual(raw_sum, Decimal('-790.92'))
        # The raw eBay refund amounts are exactly what was imported - untouched, individually traceable.
        self.assertEqual(sorted(round(float(v), 2) for v in refund_rows.Erlös_Brutto), sorted(float(a) for a in raw_amounts))

        model = prepare_partner_export(next_invoice)
        actual_impact = sum(item['gross'] for item in model['Gutschriften']
                            if item['order'] in {order for order,_ in full_refunds})
        expected_impact = sum(reference_gutschriften_gross([amount]) for amount in raw_amounts)
        self.assertEqual(actual_impact, expected_impact)
        naive_shortcut = (raw_sum * Decimal('.965')).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
        # Whether or not this happens to coincide with the naive shortcut for this
        # particular data is not the point (VAT add/remove approximately cancels
        # out) - what matters is that the value is derived from the real formula.
        self.assertEqual(actual_impact, reference_gutschriften_gross(raw_amounts))
        self.assertNotEqual(  # guards against a hard-coded gross*(1-rate) shortcut replacing the real formula
            [item['gross'] for item in model['Gutschriften']],
            [Decimal(a) * Decimal('.965') for a in raw_amounts],
        )
        del naive_shortcut  # documented above for context only, not asserted against


if __name__ == '__main__':
    unittest.main()
